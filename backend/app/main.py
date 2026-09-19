from __future__ import annotations

import json
import logging
import os
import tempfile
import time
import uuid
import zipfile
from io import BytesIO
from pathlib import Path
from typing import Any, Literal

import httpx
from fastapi import BackgroundTasks, FastAPI, File, Form, HTTPException, Query, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

from . import database

try:
    from docx import Document as WordDocument
    from docx.oxml.ns import qn
except ImportError:  # pragma: no cover
    WordDocument = None
    qn = None

try:
    from sarvamai import SarvamAI
except ImportError:  # pragma: no cover
    SarvamAI = None

# Keep the complete multipart request safely below Vercel's 4.5 MB limit.
MAX_AUDIO_BYTES = 4_000_000
MAX_AUDIO_SECONDS = 10 * 60
ALLOWED_AUDIO_TYPES = {
    "audio/webm", "audio/ogg", "audio/opus", "audio/mpeg", "audio/mp3",
    "audio/wav", "audio/x-wav", "audio/aac", "audio/flac", "audio/mp4",
    "audio/x-m4a", "video/webm", "video/mp4",
}
ALLOWED_AUDIO_EXTENSIONS = {".webm", ".ogg", ".opus", ".mp3", ".wav", ".aac", ".flac", ".m4a", ".mp4"}
TRANSIENT_STATUS_CODES = {429, 500, 503}
SARVAM_BASE_URL = "https://api.sarvam.ai"
logger = logging.getLogger("odia_voice_papers")
logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(name)s %(message)s")


class ExportRequest(BaseModel):
    format: Literal["docx", "md", "txt"]
    content: str
    title: str = "Odia document"


class ContentRequest(BaseModel):
    content: str


database.initialize()
app = FastAPI(title="Odia Voice Papers API", version="0.2.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "http://127.0.0.1:5173"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.on_event("startup")
def cleanup_old_uploads() -> None:
    removed = database.cleanup_uploads(max_age_hours=24)
    if removed:
        logger.info("removed_expired_uploads count=%s", removed)


def response_value(response: Any, name: str, default: Any = None) -> Any:
    if isinstance(response, dict):
        return response.get(name, default)
    return getattr(response, name, default)


def safe_error(error: Exception) -> str:
    message = str(error)
    api_key = os.getenv("SARVAM_API_KEY")
    if api_key:
        message = message.replace(api_key, "[redacted]")
    return message[:1000]


def normalize_result(result: Any, document_id: str) -> tuple[str, list[dict[str, Any]]]:
    transcript = response_value(result, "transcript", "") or ""
    timestamps = response_value(result, "timestamps") or {}
    chunks = response_value(timestamps, "chunks", []) or []
    starts = response_value(timestamps, "start_time_seconds", []) or []
    ends = response_value(timestamps, "end_time_seconds", []) or []
    diarized = response_value(result, "diarized_transcript") or {}
    entries = response_value(diarized, "entries", []) or []
    segments: list[dict[str, Any]] = []
    source = entries or chunks
    for index, entry in enumerate(source):
        if isinstance(entry, str):
            text, speaker_id, start, end = entry, None, starts[index] if index < len(starts) else None, ends[index] if index < len(ends) else None
        else:
            text = response_value(entry, "transcript", "")
            speaker_id = response_value(entry, "speaker_id")
            start = response_value(entry, "start_time_seconds")
            end = response_value(entry, "end_time_seconds")
        segments.append({"id": f"{document_id}-{index}", "text": text, "speaker_id": speaker_id, "start_seconds": start, "end_seconds": end})
    if not segments and transcript:
        segments.append({"id": f"{document_id}-0", "text": transcript, "speaker_id": None, "start_seconds": None, "end_seconds": None})
    return transcript, segments


def get_sarvam_client() -> Any:
    api_key = os.getenv("SARVAM_API_KEY")
    if not api_key:
        raise RuntimeError("SARVAM_API_KEY is not configured")
    if SarvamAI is None:
        raise RuntimeError("Install backend requirements before using Sarvam")
    return SarvamAI(api_subscription_key=api_key, timeout=60.0)


def download_batch_result(batch_job: Any, temp_dir: str) -> Any:
    output_dir = Path(temp_dir) / "outputs"
    output_dir.mkdir(exist_ok=True)
    batch_job.download_outputs(output_dir=str(output_dir))
    json_files = sorted(output_dir.glob("*.json"))
    if not json_files:
        raise RuntimeError("Sarvam batch completed without a JSON output file")
    return json.loads(json_files[0].read_text(encoding="utf-8"))


def submit_batch(job_id: str) -> None:
    row = database.get_job(job_id)
    if not row or not row["file_path"]:
        return
    client = get_sarvam_client()
    batch_job = client.speech_to_text_job.create_job(
        model="saaras:v3", mode="transcribe", language_code="od-IN", with_diarization=True
    )
    external_job_id = response_value(batch_job, "job_id")
    database.update_job(job_id, external_job_id=external_job_id, state="submitted", progress=35)
    batch_job.upload_files(file_paths=[row["file_path"]])
    batch_job.start()
    database.update_job(job_id, state="running", progress=50)
    database.audit("batch_submitted", row["document_id"], job_id, external_job_id=external_job_id)


def fetch_remote_batch_status(external_job_id: str) -> tuple[str, dict[str, Any], Any | None]:
    api_key = os.getenv("SARVAM_API_KEY")
    if not api_key:
        raise RuntimeError("SARVAM_API_KEY is not configured")
    headers = {"api-subscription-key": api_key}
    status_url = f"{SARVAM_BASE_URL}/speech-to-text/job/v1/{external_job_id}/status"
    with httpx.Client(timeout=30.0) as client:
        response = client.get(status_url, headers=headers)
        response.raise_for_status()
        status = response.json()
        state = str(status.get("job_state", "")).lower()
        if state in {"accepted", "pending", "running"}:
            return "running", status, None
        if state == "failed":
            return "failed", status, None
        if state != "completed":
            return state or "unknown", status, None
        output_names = [
            output.get("file_name")
            for detail in status.get("job_details", [])
            for output in detail.get("outputs", [])
            if output.get("file_name")
        ]
        if not output_names:
            raise RuntimeError("Sarvam completed without output files")
        download_url = f"{SARVAM_BASE_URL}/speech-to-text/job/v1/download-files"
        download = client.post(download_url, headers=headers, json={"job_id": external_job_id, "files": output_names})
        download.raise_for_status()
        with zipfile.ZipFile(BytesIO(download.content)) as archive:
            json_name = next(name for name in archive.namelist() if name.endswith(".json"))
            raw_result = json.loads(archive.read(json_name).decode("utf-8"))
        return "completed", status, raw_result


def poll_external_batch(job_id: str) -> None:
    row = database.get_job(job_id)
    if not row or not row["external_job_id"]:
        return
    state, status, raw_result = fetch_remote_batch_status(row["external_job_id"])
    if state == "running":
        database.update_job(job_id, state="running", progress=70)
        return
    if state == "failed":
        database.update_job(job_id, state="failed", error_message=status.get("error_message", "Sarvam batch failed"))
        return
    if state != "completed" or raw_result is None:
        return
    transcript, segments = normalize_result(raw_result, row["document_id"])
    database.save_result(job_id, row["document_id"], transcript, segments, raw_result)
    database.audit("batch_completed", row["document_id"], job_id, external_job_id=row["external_job_id"], segment_count=len(segments))


def process_batch(job_id: str) -> None:
    row = database.get_job(job_id)
    if not row:
        return
    database.update_job(job_id, state="processing", progress=25, attempts=row["attempts"] + 1)
    database.audit("batch_started", row["document_id"], job_id, attempt=row["attempts"] + 1)
    temp_dir = str(Path(row["file_path"]).parent) if row["file_path"] else ""
    try:
        client = get_sarvam_client()
        batch_job = client.speech_to_text_job.create_job(
            model="saaras:v3", mode="transcribe", language_code="od-IN", with_diarization=True
        )
        external_job_id = response_value(batch_job, "job_id")
        database.update_job(job_id, external_job_id=external_job_id, progress=35)
        batch_job.upload_files(file_paths=[row["file_path"]])
        database.update_job(job_id, progress=50)
        batch_job.start()
        batch_job.wait_until_complete(poll_interval=5, timeout=7200)
        database.update_job(job_id, progress=85)
        raw_result = download_batch_result(batch_job, temp_dir)
        transcript, segments = normalize_result(raw_result, row["document_id"])
        database.save_result(job_id, row["document_id"], transcript, segments, raw_result)
        database.audit("batch_completed", row["document_id"], job_id, external_job_id=external_job_id, segment_count=len(segments))
    except Exception as error:
        message = safe_error(error)
        status_code = getattr(error, "status_code", None)
        database.update_job(job_id, state="failed", error_message=message)
        database.audit("batch_failed", row["document_id"], job_id, status_code=status_code, message=message)
        logger.error("batch_failed job_id=%s status_code=%s message=%s", job_id, status_code, message)
    finally:
        latest = database.get_job(job_id)
        if row["file_path"] and latest and latest["state"] == "completed":
            try:
                Path(row["file_path"]).unlink(missing_ok=True)
                Path(temp_dir).rmdir()
            except OSError:
                pass


@app.get("/api/health")
def health() -> dict[str, str | int]:
    return {
        "status": "ok",
        "batch_mode": "enabled",
        "persistence": "sqlite",
        "max_audio_bytes": MAX_AUDIO_BYTES,
        "max_audio_seconds": MAX_AUDIO_SECONDS,
    }


@app.get("/api/usage")
def usage_summary() -> dict[str, Any]:
    return database.usage_summary()


@app.post("/api/documents")
async def create_document(
    background_tasks: BackgroundTasks,
    title: str = Form("Odia voice document"),
    duration_seconds: float | None = Form(None),
    file: UploadFile = File(...),
) -> dict[str, Any]:
    extension = Path(file.filename or "").suffix.lower()
    if file.content_type not in ALLOWED_AUDIO_TYPES and extension not in ALLOWED_AUDIO_EXTENSIONS:
        raise HTTPException(status_code=415, detail="Supported formats: WebM, MP3, WAV, AAC, FLAC, OGG, Opus, M4A, and MP4")
    if duration_seconds is not None and duration_seconds > MAX_AUDIO_SECONDS:
        raise HTTPException(status_code=422, detail="Audio sessions cannot exceed 10 minutes")
    data = await file.read()
    if not data:
        raise HTTPException(status_code=400, detail="The audio file is empty")
    if len(data) > MAX_AUDIO_BYTES:
        raise HTTPException(status_code=413, detail="The audio file must be smaller than 4 MB for Vercel upload limits")

    document_id, job_id = str(uuid.uuid4()), str(uuid.uuid4())
    temp_dir = tempfile.mkdtemp(prefix="odia-voice-")
    extension = extension or ".audio"
    file_path = str(Path(temp_dir) / f"source{extension}")
    Path(file_path).write_bytes(data)
    database.create_document(document_id, title)
    database.create_job(job_id, document_id, file.filename or "audio", file_path, "queued")
    database.add_usage(job_id, len(data), duration_seconds)
    database.audit("document_uploaded", document_id, job_id, bytes=len(data), duration_seconds=duration_seconds)
    if os.getenv("VERCEL"):
        try:
            submit_batch(job_id)
            Path(file_path).unlink(missing_ok=True)
            database.update_job(job_id, file_path=None)
        except Exception as error:
            database.update_job(job_id, state="failed", error_message=safe_error(error))
    else:
        background_tasks.add_task(process_batch, job_id)
    created_job = database.get_job(job_id)
    return {
        "document_id": document_id,
        "job_id": job_id,
        "state": created_job["state"] if created_job else "failed",
        "external_job_id": created_job["external_job_id"] if created_job else None,
    }


@app.get("/api/transcription-jobs/{job_id}")
def get_job(
    job_id: str,
    external_job_id: str | None = Query(None),
    document_id: str | None = Query(None),
) -> dict[str, Any]:
    row = database.get_job(job_id)
    if not row:
        if not external_job_id:
            raise HTTPException(status_code=404, detail="Job not found")
        state, status, raw_result = fetch_remote_batch_status(external_job_id)
        transcript, segments = ("", [])
        if state == "completed" and raw_result is not None:
            transcript, segments = normalize_result(raw_result, document_id or job_id)
        return {
            "state": state,
            "progress": 100 if state == "completed" else 70,
            "transcript": transcript,
            "segments": segments,
            "error_message": status.get("error_message") if state == "failed" else None,
            "external_job_id": external_job_id,
        }
    if os.getenv("VERCEL") and row["external_job_id"]:
        try:
            poll_external_batch(job_id)
            row = database.get_job(job_id) or row
        except Exception as error:
            message = safe_error(error)
            database.update_job(job_id, state="failed", error_message=message)
            database.audit("batch_poll_failed", row["document_id"], job_id, message=message)
            logger.error("client_poll_failed job_id=%s message=%s", job_id, message)
            row = database.get_job(job_id) or row
    return {
        "state": row["state"], "progress": row["progress"], "transcript": database.get_document_content(row["document_id"]),
        "segments": database.get_segments(job_id), "error_message": row["error_message"], "external_job_id": row["external_job_id"],
    }


@app.post("/api/transcription-jobs/{job_id}/retry")
def retry_job(job_id: str, background_tasks: BackgroundTasks) -> dict[str, str]:
    row = database.get_job(job_id)
    if not row:
        raise HTTPException(status_code=404, detail="Job not found")
    if row["state"] != "failed":
        raise HTTPException(status_code=409, detail="Only failed jobs can be retried")
    if row["attempts"] >= 3:
        raise HTTPException(status_code=429, detail="Retry limit reached")
    database.update_job(job_id, state="queued", progress=15, error_message=None)
    database.audit("batch_retry_requested", row["document_id"], job_id, attempt=row["attempts"] + 1)
    if os.getenv("VERCEL"):
        try:
            submit_batch(job_id)
            Path(row["file_path"]).unlink(missing_ok=True)
            database.update_job(job_id, file_path=None)
        except Exception as error:
            database.update_job(job_id, state="failed", error_message=safe_error(error))
    else:
        background_tasks.add_task(process_batch, job_id)
    return {"job_id": job_id, "state": "queued"}


@app.get("/api/documents/{document_id}")
def get_document(document_id: str) -> dict[str, Any]:
    document = database.get_document(document_id)
    if not document:
        raise HTTPException(status_code=404, detail="Document not found")
    return {"id": document["id"], "title": document["title"], "status": document["status"], "content": database.get_document_content(document_id), "updated_at": document["updated_at"]}


@app.patch("/api/documents/{document_id}/content")
def save_document_content(document_id: str, request: ContentRequest) -> dict[str, Any]:
    if not database.get_document(document_id):
        raise HTTPException(status_code=404, detail="Document not found")
    version = database.save_edited_version(document_id, request.content)
    database.audit("document_edited", document_id, version=version, characters=len(request.content))
    return {"document_id": document_id, "version": version, "saved": True}


@app.get("/api/documents/{document_id}/versions")
def document_versions(document_id: str) -> dict[str, Any]:
    if not database.get_document(document_id):
        raise HTTPException(status_code=404, detail="Document not found")
    return {"versions": database.get_versions(document_id)}


@app.get("/api/audit")
def audit_events(limit: int = 100) -> dict[str, Any]:
    return {"events": database.get_audit_events(min(max(limit, 1), 500))}


@app.post("/api/documents/{document_id}/exports")
def export_document(document_id: str, request: ExportRequest) -> StreamingResponse:
    if not database.get_document(document_id):
        raise HTTPException(status_code=404, detail="Document not found")
    safe_title = "".join(character for character in request.title if character.isalnum() or character in " _-") or "odia-document"
    database.audit("document_exported", document_id, format=request.format)
    if request.format == "docx":
        if WordDocument is None:
            raise HTTPException(status_code=500, detail="python-docx is not installed")
        document = WordDocument()
        for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
            style = document.styles[style_name]
            style.font.name = "Kalinga"
            style._element.rPr.rFonts.set(qn("w:ascii"), "Kalinga")
            style._element.rPr.rFonts.set(qn("w:hAnsi"), "Kalinga")
            style._element.rPr.rFonts.set(qn("w:eastAsia"), "Kalinga")
        heading = document.add_heading(request.title, level=1)
        for run in heading.runs:
            run.font.name = "Kalinga"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Kalinga")
        for paragraph in request.content.split("\n"):
            paragraph_node = document.add_paragraph(paragraph)
            for paragraph_run in paragraph_node.runs:
                paragraph_run.font.name = "Kalinga"
                paragraph_run._element.rPr.rFonts.set(qn("w:eastAsia"), "Kalinga")
        output = BytesIO()
        document.save(output)
        output.seek(0)
        return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={"Content-Disposition": f'attachment; filename="{safe_title}.docx"'})
    media_type = "text/markdown" if request.format == "md" else "text/plain"
    return StreamingResponse(iter([request.content.encode("utf-8")]), media_type=f"{media_type}; charset=utf-8", headers={"Content-Disposition": f'attachment; filename="{safe_title}.{request.format}"'})


frontend_directory = Path(__file__).resolve().parents[2] / "frontend"
if frontend_directory.exists():
    app.mount("/", StaticFiles(directory=frontend_directory, html=True), name="frontend")
